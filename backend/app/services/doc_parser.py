from pathlib import Path


def parse_document(file_path: str, file_type: str) -> str:
    path = Path(file_path)
    suffix = file_type or path.suffix.lower().lstrip(".")

    if suffix == "pdf":
        return _parse_pdf(path)
    if suffix in ("docx", "doc"):
        return _parse_docx(path)
    if suffix in ("md", "markdown"):
        return _parse_text(path)
    if suffix in ("txt", "text", "log", "csv"):
        return _parse_text(path)
    raise ValueError(f"unsupported file type: {suffix}")


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text.strip())
    return "\n\n".join(p for p in pages if p)


def _parse_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_text(path: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return text.strip()


def split_chunks(text: str, chunk_size: int = 500, overlap: int = 60) -> list[str]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", "。", "！", "？", "；", ".", "!", "?", ";", " ", ""],
    )
    chunks = [c.strip() for c in splitter.split_text(text) if c.strip()]
    return chunks
